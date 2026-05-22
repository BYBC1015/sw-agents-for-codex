#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image


FORBIDDEN_TEXT = [
    "Midjourney Prompt",
    "Midjourney 提示词",
    "--ar ",
    "正选图",
    "场次 / 正选图 / 备注",
    "生成结果",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9DEE7", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_fixed_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            row.cells[idx].width = Inches(width)


def add_run(paragraph, text: str, *, bold=False, color: str | None = None):
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in [
        ("Title", 24, "111827", 0, 8),
        ("Subtitle", 11, "4B5563", 0, 14),
        ("Heading 1", 16, "1F4D78", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 11, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_key_value_table(doc: Document, rows: list[list[str] | tuple[str, str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_borders(table, color="D9DEE7", size="4")
    set_fixed_table_width(table, [1.25, 5.0])
    for idx, pair in enumerate(rows):
        key, value = pair[0], pair[1]
        cells = table.rows[idx].cells
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        set_cell_shading(cells[0], "F2F4F7")
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        add_run(p0, str(key), bold=True, color="1F2937")
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        add_run(p1, str(value))


def add_picture_fit(doc: Document, image_path: Path, max_width_in=6.15) -> None:
    if not image_path.exists():
        raise FileNotFoundError(f"Image not found: {image_path}")
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    if width_px <= 0 or height_px <= 0:
        raise ValueError(f"Invalid image dimensions: {image_path}")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(max_width_in))


def add_scene_card(doc: Document, scene: dict) -> None:
    doc.add_heading(scene["title"], level=2)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(8)
    add_run(meta, scene.get("range", ""), bold=True, color="4B5563")

    add_picture_fit(doc, Path(scene["image"]).expanduser(), max_width_in=float(scene.get("image_width_in", 6.15)))

    add_key_value_table(
        doc,
        [
            ("叙事功能", scene.get("function", "")),
            ("色调/光线", scene.get("tone", "")),
            ("视觉符号", scene.get("symbols", "")),
        ],
    )

    setting = scene.get("setting", "")
    if setting:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        add_run(p, "视觉设定：", bold=True, color="1F4D78")
        add_run(p, setting)

    reason = scene.get("reason", "")
    if reason:
        p = doc.add_paragraph()
        add_run(p, "为什么这样设定：", bold=True, color="1F4D78")
        add_run(p, reason)


def validate_spec(spec: dict) -> None:
    scenes = spec.get("scenes") or []
    if not scenes:
        raise ValueError("Spec must include at least one scene.")
    for i, scene in enumerate(scenes, 1):
        for key in ("title", "image", "setting", "reason"):
            if not scene.get(key):
                raise ValueError(f"Scene {i} missing required key: {key}")
        if any(token in json.dumps(scene, ensure_ascii=False) for token in FORBIDDEN_TEXT[:3]):
            raise ValueError(f"Scene {i} appears to contain prompt text; remove prompts from the Word spec.")


def build_docx(spec: dict, output_path: Path) -> None:
    validate_spec(spec)
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(spec.get("title", "视觉导演方案"))

    subtitle_text = spec.get("subtitle")
    if subtitle_text:
        subtitle = doc.add_paragraph(style="Subtitle")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.add_run(subtitle_text)

    add_key_value_table(doc, spec.get("metadata", []))

    scenes = spec["scenes"]
    doc.add_heading(spec.get("scene_split_heading", "场景拆分"), level=1)
    for scene in scenes:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        add_run(paragraph, f"{scene['title']}：", bold=True, color="1F4D78")
        add_run(paragraph, scene.get("range", ""))

    doc.add_heading(spec.get("scene_text_heading", "视觉设定成片文本"), level=1)
    page_break_before = set(spec.get("page_break_before_scene_indices", []))
    if not page_break_before and len(scenes) >= 6:
        page_break_before = {3, 5, 7}
    for idx, scene in enumerate(scenes, 1):
        if idx in page_break_before:
            doc.add_page_break()
        add_scene_card(doc, scene)

    qa_items = spec.get("qa", [])
    if qa_items:
        doc.add_heading(spec.get("qa_heading", "QA 复盘"), level=1)
        for item in qa_items:
            doc.add_paragraph(str(item), style="List Bullet")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def inspect_docx(path: Path) -> dict:
    with ZipFile(path) as zf:
        media = [name for name in zf.namelist() if name.startswith("word/media/")]
        document_xml = zf.read("word/document.xml")
        rels = zf.read("word/_rels/document.xml.rels")
    root = etree.fromstring(document_xml)
    text = "".join(root.xpath("//w:t/text()", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}))
    forbidden_found = [token for token in FORBIDDEN_TEXT if token in text]
    return {
        "docx": str(path),
        "size_bytes": path.stat().st_size,
        "embedded_media_count": len(media),
        "image_relationship_count": rels.count(b'Target="media/'),
        "forbidden_text_found": forbidden_found,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a clean storyboard visual-director DOCX with embedded scene images.")
    parser.add_argument("--spec", required=True, help="JSON spec containing scenes and image paths.")
    parser.add_argument("--out", required=True, help="Output .docx path.")
    parser.add_argument("--inspect-json", help="Optional path to write validation/inspection JSON.")
    args = parser.parse_args()

    spec_path = Path(args.spec).expanduser()
    output_path = Path(args.out).expanduser()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    build_docx(spec, output_path)
    report = inspect_docx(output_path)
    if report["forbidden_text_found"]:
        raise RuntimeError(f"Forbidden visible text found in DOCX: {report['forbidden_text_found']}")
    if report["embedded_media_count"] < len(spec["scenes"]):
        raise RuntimeError("DOCX embedded fewer images than scenes.")
    if args.inspect_json:
        Path(args.inspect_json).expanduser().write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
